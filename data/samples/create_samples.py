"""Script to generate realistic academic sample files for .docx and .pdf formats."""

from pathlib import Path


def create_docx(dest: Path):
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    title = doc.add_heading("Evaluating Integrity in Generative Text Systems", level=0)

    doc.add_heading("Abstract", level=1)
    p_abs = doc.add_paragraph(
        "Automated plagiarism and integrity verification tools must operate robustly across diverse file "
        "formats. This document serves as a standard test sample for Word document ingestion, paragraph extraction, "
        "and character offset tracking."
    )

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Modern academic environments demand transparent mechanisms to verify student submissions. "
        "Textual comparison algorithms analyze lexical, syntactic, and semantic signatures."
    )

    doc.add_heading("2. Theoretical Principles", level=1)
    doc.add_paragraph(
        "Direct quotation is a fundamental scholarly practice, provided attribution is explicit: "
        '"Scholarly inquiry thrives when ideas are openly cited and acknowledged across disciplines."'
    )

    doc.add_heading("3. Methodology", level=1)
    doc.add_paragraph(
        "Documents are parsed into canonical plain text streams before sentence boundary tokenization. "
        "Spans are computed relative to the cleaned text representation to facilitate highlighting."
    )

    doc.add_heading("References", level=1)
    doc.add_paragraph("1. Manning, C. D., Raghavan, P., & Schütze, H. (2008). Introduction to Information Retrieval. Cambridge University Press.")
    doc.add_paragraph("2. Jurafsky, D., & Martin, J. H. (2023). Speech and Language Processing (3rd ed.).")

    doc.save(dest)
    print(f"Generated {dest}")


def create_pdf(dest: Path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Title
    pdf.set_font("Helvetica", style="B", size=16)
    pdf.cell(0, 10, "Heuristic Approaches to Document Verification", ln=True, align="C")
    pdf.ln(5)

    # Abstract
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "Abstract", ln=True)
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(
        0,
        5,
        "Ensuring the authenticity of written assignments requires dependable document ingestion. "
        "This PDF file serves as a test artifact to validate text extraction from portable document formats, "
        "verifying that layout streams are converted into contiguous plain text.",
    )
    pdf.ln(4)

    # Introduction
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "1. Introduction", ln=True)
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(
        0,
        5,
        "Document ingestion forms the primary stage of text analysis pipelines. "
        "Discrepancies in whitespace, ligature decoding, and hyphenation must be normalized. "
        "As stated by modern linguists, \"Precise text normalization prevents cascading errors in downstream semantic models.\"",
    )
    pdf.ln(4)

    # Methodology
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "2. Processing Methodology", ln=True)
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(
        0,
        5,
        "The ingestion engine parses page streams sequentially. "
        "Headers, footers, and reference lists are categorized to isolate the substantive argumentative prose.",
    )
    pdf.ln(4)

    # References
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "References", ln=True)
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(
        0,
        5,
        "1. Bird, S., Klein, E., & Loper, E. (2009). Natural Language Processing with Python. O'Reilly Media.\n"
        "2. Russell, S., & Norvig, P. (2020). Artificial Intelligence: A Modern Approach (4th ed.). Pearson.",
    )

    pdf.output(str(dest))
    print(f"Generated {dest}")


if __name__ == "__main__":
    base_dir = Path(__file__).resolve().parent
    create_docx(base_dir / "sample_academic.docx")
    create_pdf(base_dir / "sample_academic.pdf")
