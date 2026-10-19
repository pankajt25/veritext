"""Document loaders for plain text extraction across .docx, .pdf, .tex, .txt, and .md files."""

from pathlib import Path
from typing import Union
from veritext.ingest.latex import process_latex


def read_text_file(path: Union[str, Path]) -> str:
    """Reads a text file handling UTF-8 and fallback encodings."""
    path = Path(path)
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")


def load_txt(path: Union[str, Path]) -> str:
    """Loads plain text from a .txt file."""
    return read_text_file(path)


def load_md(path: Union[str, Path]) -> str:
    """Loads text from a Markdown (.md) file."""
    return read_text_file(path)


def load_tex(path: Union[str, Path]) -> str:
    """Loads and converts a LaTeX (.tex) document into plain text."""
    raw_tex = read_text_file(path)
    plain_text, _ = process_latex(raw_tex)
    return plain_text


def load_docx(path: Union[str, Path]) -> str:
    """Loads plain text from a Microsoft Word (.docx) document."""
    path = Path(path)
    import docx

    doc = docx.Document(path)
    parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Also capture table contents if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n\n".join(parts)


def load_pdf(path: Union[str, Path]) -> str:
    """Loads plain text from a PDF (.pdf) file using pdfplumber, falling back to OCR if text is empty."""
    path = Path(path)
    import pdfplumber

    page_texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text and text.strip():
                page_texts.append(text.strip())

    extracted = "\n\n".join(page_texts).strip()
    if not extracted:
        # Fallback to OCR for scanned PDF
        from veritext.ingest.ocr import ocr_scanned_pdf
        extracted = ocr_scanned_pdf(path)

    return extracted


def load_image(path: Union[str, Path]) -> str:
    """Loads text from an image file (.png, .jpg, .jpeg) using OCR."""
    from veritext.ingest.ocr import ocr_image
    return ocr_image(path)


LOADER_MAP = {
    ".txt": load_txt,
    ".md": load_md,
    ".markdown": load_md,
    ".tex": load_tex,
    ".latex": load_tex,
    ".docx": load_docx,
    ".pdf": load_pdf,
    ".png": load_image,
    ".jpg": load_image,
    ".jpeg": load_image,
}


def get_loader(path: Union[str, Path]):
    """Returns the appropriate loader function based on file extension."""
    suffix = Path(path).suffix.lower()
    if suffix not in LOADER_MAP:
        raise ValueError(
            f"Unsupported file format '{suffix}'. Supported formats: {', '.join(LOADER_MAP.keys())}"
        )
    return LOADER_MAP[suffix]
